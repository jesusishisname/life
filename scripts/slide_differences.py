import io
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

import fitz
import requests
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.util import Inches

url = os.environ["SLIDES_URL"].strip()
match = re.search(r"/presentation/d/([a-zA-Z0-9_-]+)", url)
if not match or urlparse(url).netloc not in {"docs.google.com", "slides.google.com"}:
    raise SystemExit("Enter a public Google Slides presentation URL.")

presentation_id = match.group(1)
pdf_url = f"https://docs.google.com/presentation/d/{presentation_id}/export/pdf"
response = requests.get(pdf_url, timeout=90, allow_redirects=True)
response.raise_for_status()
if not response.content.startswith(b"%PDF"):
    raise SystemExit("Could not export this presentation as PDF. Check that link sharing allows anyone with the link to view it.")

work = Path("work")
work.mkdir(exist_ok=True)
doc = fitz.open(stream=response.content, filetype="pdf")
if len(doc) < 2:
    raise SystemExit("The presentation must contain at least two slides.")

images = []
for i, page in enumerate(doc):
    # Save each slide as SVG and render it for pixel-wise differencing.
    (work / f"slide-{i + 1:03}.svg").write_text(page.get_svg_image(), encoding="utf-8")
    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
    images.append(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB"))

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for i in range(len(images) - 1):
    a, b = images[i], images[i + 1]
    if b.size != a.size:
        b = b.resize(a.size, Image.Resampling.LANCZOS)
    diff = ImageChops.difference(b, a)
    # Enhance subtle pixel changes so they remain visible in the result.
    diff = diff.point(lambda p: min(255, p * 3))
    path = work / f"difference-{i + 1:03}-to-{i + 2:03}.png"
    diff.save(path)
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)

out = Path("outputs")
out.mkdir(exist_ok=True)
pptx_path = out / "slide-differences.pptx"
prs.save(pptx_path)
print(f"Created {pptx_path} with {len(images) - 1} difference slides.")
